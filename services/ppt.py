"""PPT 生成：LLM 生成幻灯片内容 -> python-pptx 生成 .pptx。

内置 4 套主题模板（配色 + 真实背景图），供 getppt 选择、generate_ppt 渲染。
每页 PNG 预览由服务端 LibreOffice 从 .pptx 渲染；本模块仅用 Pillow 渲染主题缩略图。
"""
import base64
import io
import json
import os
import random
import re
import urllib.request

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from services import image_search
from services import llm

_FONT_NAME = "微软雅黑"

_TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "templates")

def _load_themes():
    """从 templates/themes.json 加载主题配置，颜色转 tuple，并加载 mode.json 版式参数。"""
    path = os.path.join(_TEMPLATE_DIR, "themes.json")
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)
    themes = []
    for t in raw:
        t = dict(t)
        for key in ("accent", "title_color", "body_color", "footer_color", "cover_bg"):
            t[key] = tuple(t[key])
        mode_file = t.get("mode")
        if mode_file:
            mode_path = os.path.join(_TEMPLATE_DIR, mode_file)
            try:
                with open(mode_path, "r", encoding="utf-8") as mf:
                    t["mode"] = json.load(mf)
            except Exception:
                t["mode"] = {}
        else:
            t["mode"] = {}
        themes.append(t)
    return themes


# 内置主题模板：id 与前端选中值一致，name 用于展示
# bg 为背景图文件名（相对 templates 目录），cover_bg 为无背景图时的兜底色
THEMES = _load_themes()


def _set_run_font(run):
    """给 run 设置中文字体（latin 与 east asian 都设，否则中文吃不到字体）。"""
    run.font.name = _FONT_NAME
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", _FONT_NAME)


def _bg_path(theme):
    """返回内容页随机背景图路径；无 backgrounds 时回退 bg（兼容旧配置）。"""
    backgrounds = theme.get("backgrounds")
    name = random.choice(backgrounds) if backgrounds else theme.get("bg")
    if not name:
        return None
    path = os.path.join(_TEMPLATE_DIR, name)
    return path if os.path.exists(path) else None


def _title_image_path(theme):
    """返回封面 title 图路径；无 title_image 时回退 bg。"""
    name = theme.get("title_image") or theme.get("bg")
    if not name:
        return None
    path = os.path.join(_TEMPLATE_DIR, name)
    return path if os.path.exists(path) else None


def _cm2in(cm):
    return cm / 2.54


def _mode_box(theme, page_type, element):
    """从 mode.json 取某页型某元素的 (left, top, width, height)，单位英寸；缺失返回 None。"""
    info = theme.get("mode", {}).get(page_type, {}).get(element)
    if not info:
        return None
    return (
        _cm2in(info.get("pos_x", 0)),
        _cm2in(info.get("pos_y", 0)),
        _cm2in(info.get("width", 0)),
        _cm2in(info.get("height", 0)),
    )


def _mode_font(theme, page_type, element, default=18):
    """从 mode.json 取某页型某元素字号，缺失返回 default。"""
    info = theme.get("mode", {}).get(page_type, {}).get(element, {})
    return info.get("font_size", default)


def _slide_size(theme):
    """从 mode.json 取幻灯片尺寸 (width_in, height_in)，缺失回退 16:9 13.333x7.5。"""
    size = theme.get("mode", {}).get("slide_size", {})
    w, h = size.get("width"), size.get("height")
    if w and h:
        return _cm2in(w), _cm2in(h)
    return 13.333, 7.5


def _add_page_background(slide, theme, prs):
    """给内容页铺随机背景图；无图时用白底。"""
    bg = _bg_path(theme)
    if bg:
        slide.shapes.add_picture(bg, 0, 0, width=prs.slide_width, height=prs.slide_height)
    else:
        _set_bg(slide, (255, 255, 255))


def _fit(img: Image.Image, width: int, height: int) -> Image.Image:
    """cover 模式裁剪到目标比例后缩放。"""
    ratio = width / height
    iratio = img.width / img.height
    if iratio > ratio:
        new_w = int(img.height * ratio)
        left = (img.width - new_w) // 2
        img = img.crop((left, 0, left + new_w, img.height))
    else:
        new_h = int(img.width / ratio)
        top = (img.height - new_h) // 2
        img = img.crop((0, top, img.width, top + new_h))
    return img.resize((width, height), Image.LANCZOS)


def _extract_json(raw: str):
    """从 LLM 输出里提取 JSON 数组（容忍 markdown 代码块/前后缀）。"""
    raw = (raw or "").strip()
    if raw.startswith("```"):
        lines = raw.split("\n")
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        raw = "\n".join(lines).strip()
    m = re.search(r"\[.*\]", raw, re.DOTALL)
    if m:
        raw = m.group(0)
    if not raw:
        raise ValueError("LLM 返回内容为空，无法生成幻灯片")
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        raise ValueError(f"LLM 返回不是合法 JSON: {raw[:200]}")


_LAYOUTS = ("cover", "agenda", "section", "content", "image_text", "full_image", "quote", "comparison", "closing")

_LAYOUT_SCHEMA = (
    "版式字段规范（layout 九选一）：\n"
    'cover 封面: {"layout":"cover","title":"主标题","subtitle":"副标题","notes":"备注"}\n'
    'agenda 目录: {"layout":"agenda","items":["章节1","章节2"]}\n'
    'section 章节过渡: {"layout":"section","title":"章节标题"}\n'
    'content 标题要点: {"layout":"content","title":"页标题","bullets":["要点1"],"notes":"备注"}\n'
    'image_text 左图右文: {"layout":"image_text","title":"...","bullets":["..."],"image":"配图关键词","image_side":"left"}\n'
    'full_image 全图: {"layout":"full_image","title":"...","image":"配图关键词"}\n'
    'quote 金句: {"layout":"quote","text":"引文","source":"出处","image":"配图关键词"}\n'
    'comparison 对比: {"layout":"comparison","title":"...","left":{"title":"A","points":["..."]},"right":{"title":"B","points":["..."]}}\n'
    'closing 结束: {"layout":"closing","title":"谢谢/总结"}\n'
)


def generate_slides(title: str, body_text: str, system=None):
    """LLM 生成幻灯片内容，返回带 layout 字段的 9 版式结构化 JSON 数组。"""
    prompt = (
        "请把下面的课件文案整理成 PPT 幻灯片内容，只输出一个 JSON 数组（不要输出代码块或任何解释文字）。\n"
        "每个元素是一个幻灯片对象，含 layout 字段及各版式专属字段。\n"
        + _LAYOUT_SCHEMA +
        "规则：第一页用 cover，最后一页用 closing；正文内容页用 content，需要配图/强调用 image_text、full_image 或 quote，需要对比用 comparison。\n"
        "控制在 6~10 页。\n"
        f"标题：{title}\n"
        f"文案正文：\n{body_text}\n"
    )
    raw = llm.chat([{"role": "user", "content": prompt}], temperature=0.3, max_tokens=16384, system=system)
    return _normalize_slides(_extract_json(raw))


def _normalize_slides(slides):
    """规则兜底：首项强制 cover、末项强制 closing、未知/缺省 layout 回退 content。"""
    if not isinstance(slides, list) or not slides:
        raise ValueError("LLM 返回内容为空，无法生成幻灯片")
    out = []
    n = len(slides)
    for i, item in enumerate(slides):
        if not isinstance(item, dict):
            item = {"title": str(item)}
        item = dict(item)
        layout = str(item.get("layout", "")).strip()
        if layout not in _LAYOUTS:
            layout = "content"
        if i == 0:
            layout = "cover"
        elif i == n - 1:
            layout = "closing"
        item["layout"] = layout
        out.append(item)
    return out


def _resolve_slide_image(item, layout):
    """取该页配图字节；无关键词或搜图失败返回 None。"""
    keyword = str(item.get("image", "")).strip()
    if not keyword:
        return None
    if layout in ("full_image", "quote"):
        keyword = keyword + " 壁纸"
    return image_search.search(keyword)


def _add_image_placeholder(slide, x=0.9, y=1.8, w=5.2, h=4.9):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(235, 238, 243)
    box.line.color.rgb = RGBColor(210, 214, 220)
    return box


def get_template(template_id=None):
    """按 id 或名称取主题，找不到时回退到第一套。"""
    if template_id is None or template_id == "":
        return THEMES[0]
    key = str(template_id)
    for t in THEMES:
        if str(t["id"]) == key or t["name"] == template_id:
            return t
    return THEMES[0]


def list_templates():
    """返回主题列表（含缩略图 base64），供 getppt 使用。"""
    result = []
    for t in THEMES:
        png = _render_thumbnail(t)
        b64 = base64.b64encode(png).decode("ascii")
        result.append(
            {
                "id": t["id"],
                "ppt_name": t["name"],
                "subject": t["name"],
                "image": f"data:image/png;base64,{b64}",
            }
        )
    return result


def _render_thumbnail(theme, width: int = 320, height: int = 180) -> bytes:
    """生成主题缩略图（用封面 title 图）。"""
    path = _title_image_path(theme)
    if path:
        try:
            img = _fit(Image.open(path).convert("RGB"), width, height)
        except Exception:
            img = Image.new("RGB", (width, height), theme["cover_bg"])
    else:
        img = Image.new("RGB", (width, height), theme["cover_bg"])
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def build_pptx(slides, theme=None) -> bytes:
    """根据幻灯片内容（9 版式）+ 主题生成 .pptx 字节。"""
    theme = theme or THEMES[0]
    prs = Presentation()
    slide_w, slide_h = _slide_size(theme)
    prs.slide_width = Inches(slide_w)
    prs.slide_height = Inches(slide_h)
    blank = prs.slide_layouts[6]
    total = len(slides)
    for i, item in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        _style_slide(slide, item, theme, i, total, prs)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _style_slide(slide, item, theme, index, total, prs):
    notes = str(item.get("notes", ""))
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    layout = item.get("layout") or "content"
    {
        "cover": _style_cover,
        "agenda": _style_agenda,
        "section": _style_section,
        "content": _style_content,
        "image_text": _style_image_text,
        "full_image": _style_full_image,
        "quote": _style_quote,
        "comparison": _style_comparison,
        "closing": _style_closing,
    }[layout](slide, item, theme, index, total, prs)


def _set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*color)


def _add_textbox(slide, x, y, w, h, text, theme, size=18, color=None, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text or "").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*(color if color is not None else theme["body_color"]))
        _set_run_font(run)
    return tb


def _add_bullets(slide, x, y, w, h, bullets, theme, size=18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for j, b in enumerate(bullets or []):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        run = p.add_run()
        run.text = "• " + str(b)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(*theme["body_color"])
        _set_run_font(run)
    return tb


def _add_footer(slide, theme, index, total, prs):
    sw = prs.slide_width / 914400.0
    sh = prs.slide_height / 914400.0
    _add_textbox(slide, sw - 2.5, sh - 0.5, 2.3, 0.4, f"{index + 1} / {total}",
                 theme, size=10, color=theme["footer_color"], align=PP_ALIGN.RIGHT)


def _style_cover(slide, item, theme, index, total, prs):
    bg = _title_image_path(theme)
    if bg:
        slide.shapes.add_picture(bg, 0, 0, width=prs.slide_width, height=prs.slide_height)
    else:
        _set_bg(slide, theme["cover_bg"])
    box = _mode_box(theme, "first_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "")), theme,
                     size=_mode_font(theme, "first_page", "title_info", 36), color=(255, 255, 255), bold=True)


def _style_agenda(slide, item, theme, index, total, prs):
    _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "catalog_page", "title_info")
    if box:
        _add_textbox(slide, *box, "目录", theme,
                     size=_mode_font(theme, "catalog_page", "title_info", 40), color=theme["title_color"], bold=True)
    tx = box[0] if box else 0.5
    ty = (box[1] + box[3] + 0.5) if box else 1.5
    for j, it in enumerate(item.get("items", []) or []):
        _add_textbox(slide, tx, ty + j * 0.7, 8, 0.6, f"{j + 1}. {it}", theme, size=20, color=theme["body_color"])
    _add_footer(slide, theme, index, total, prs)


def _style_section(slide, item, theme, index, total, prs):
    _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "first_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "")), theme,
                     size=_mode_font(theme, "first_page", "title_info", 36), color=(255, 255, 255), bold=True)


def _style_content(slide, item, theme, index, total, prs):
    _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "main_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "")), theme,
                     size=_mode_font(theme, "main_page", "title_info", 28), color=theme["title_color"], bold=True)
    box = _mode_box(theme, "main_page", "content_info")
    if box:
        _add_bullets(slide, *box, [str(b) for b in item.get("bullets", [])], theme,
                     size=_mode_font(theme, "main_page", "content_info", 18))
    _add_footer(slide, theme, index, total, prs)


def _style_image_text(slide, item, theme, index, total, prs):
    _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "main_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "")), theme,
                     size=_mode_font(theme, "main_page", "title_info", 28), color=theme["title_color"], bold=True)
    img_box = _mode_box(theme, "main_page", "img_info")
    content_box = _mode_box(theme, "main_page", "content_info")
    data = _resolve_slide_image(item, "image_text")
    if content_box:
        if img_box:
            # 模版有 img_info：图放 img_box，文字放 content_box
            _add_bullets(slide, *content_box, [str(b) for b in item.get("bullets", [])], theme,
                         size=_mode_font(theme, "main_page", "content_info", 18))
        else:
            # 模版无 img_info：图左文右，各占 content 区一半
            half = content_box[2] / 2
            _add_bullets(slide, content_box[0] + half, content_box[1], half, content_box[3],
                         [str(b) for b in item.get("bullets", [])], theme,
                         size=_mode_font(theme, "main_page", "content_info", 18))
            img_box = (content_box[0], content_box[1], half, content_box[3])
    if img_box:
        if data is None:
            _add_image_placeholder(slide, *img_box)
        else:
            try:
                slide.shapes.add_picture(io.BytesIO(data), Inches(img_box[0]), Inches(img_box[1]),
                                         Inches(img_box[2]), Inches(img_box[3]))
            except Exception:
                _add_image_placeholder(slide, *img_box)
    _add_footer(slide, theme, index, total, prs)


def _style_full_image(slide, item, theme, index, total, prs):
    data = _resolve_slide_image(item, "full_image")
    if data is None:
        _add_page_background(slide, theme, prs)
    else:
        try:
            slide.shapes.add_picture(io.BytesIO(data), 0, 0, width=prs.slide_width, height=prs.slide_height)
        except Exception:
            _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "main_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "")), theme,
                     size=_mode_font(theme, "main_page", "title_info", 28), color=(255, 255, 255), bold=True)
    _add_footer(slide, theme, index, total, prs)


def _style_quote(slide, item, theme, index, total, prs):
    data = _resolve_slide_image(item, "quote")
    if data is not None:
        try:
            slide.shapes.add_picture(io.BytesIO(data), 0, 0, width=prs.slide_width, height=prs.slide_height)
        except Exception:
            data = None
    if data is None:
        _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "main_page", "content_info")
    if box:
        _add_textbox(slide, *box, str(item.get("text", "")), theme,
                     size=_mode_font(theme, "main_page", "content_info", 28),
                     color=(255, 255, 255) if data is not None else theme["title_color"], bold=True)
    source = item.get("source")
    if source:
        tx = box[0] if box else 0.5
        ty = (box[1] + box[3] + 0.3) if box else 4.0
        _add_textbox(slide, tx, ty, 8, 0.5, "—— " + str(source), theme, size=16,
                     color=(240, 240, 240) if data is not None else theme["footer_color"])
    _add_footer(slide, theme, index, total, prs)


def _style_comparison(slide, item, theme, index, total, prs):
    _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "main_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "")), theme,
                     size=_mode_font(theme, "main_page", "title_info", 28), color=theme["title_color"], bold=True)
    left = item.get("left") or {}
    right = item.get("right") or {}
    cb = _mode_box(theme, "main_page", "content_info")
    if cb:
        half = cb[2] / 2
        _add_textbox(slide, cb[0], cb[1], half, 0.6, str(left.get("title", "")), theme, size=22, color=theme["title_color"], bold=True)
        _add_bullets(slide, cb[0], cb[1] + 0.7, half, cb[3] - 0.7, [str(p) for p in left.get("points", [])], theme, size=16)
        _add_textbox(slide, cb[0] + half, cb[1], half, 0.6, str(right.get("title", "")), theme, size=22, color=theme["title_color"], bold=True)
        _add_bullets(slide, cb[0] + half, cb[1] + 0.7, half, cb[3] - 0.7, [str(p) for p in right.get("points", [])], theme, size=16)
    _add_footer(slide, theme, index, total, prs)


def _style_closing(slide, item, theme, index, total, prs):
    _add_page_background(slide, theme, prs)
    box = _mode_box(theme, "first_page", "title_info")
    if box:
        _add_textbox(slide, *box, str(item.get("title", "") or "谢谢/总结"), theme,
                     size=_mode_font(theme, "first_page", "title_info", 36), color=(255, 255, 255), bold=True)


# ============ fabric 画布 JSON -> .pptx ============

_PX2EMU = 9525  # 1px = 9525 EMU（1280px 对应 13.333 英寸）


def _hex_to_rgb(h):
    h = (h or "#ffffff").lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    try:
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return (255, 255, 255)


def build_pptx_from_elements(slides_elements, theme=None) -> bytes:
    """根据前端 fabric 画布 JSON（每页元素）生成 .pptx。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for raw in slides_elements:
        data = raw
        if isinstance(raw, str):
            try:
                data = json.loads(raw)
            except Exception:
                data = {}
        slide = prs.slides.add_slide(blank)
        for obj in (data or {}).get("objects", []):
            otype = obj.get("type")
            if otype == "rect" and obj.get("id") == "ppt-bg":
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(*_hex_to_rgb(obj.get("fill")))
            elif otype == "textbox":
                _add_textbox_shape(slide, obj)
            elif otype == "image":
                _add_image_shape(slide, obj)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_textbox_shape(slide, obj):
    left = int(obj.get("left", 0) * _PX2EMU)
    top = int(obj.get("top", 0) * _PX2EMU)
    width = int((obj.get("width", 300) or 0) * (obj.get("scaleX", 1) or 1) * _PX2EMU)
    height = int((obj.get("height", 100) or 0) * (obj.get("scaleY", 1) or 1) * _PX2EMU)
    width = max(width, _PX2EMU)
    height = max(height, _PX2EMU)
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    text = obj.get("text", "") or ""
    lines = text.split("\n")
    fs = int((obj.get("fontSize", 24) or 24) * 0.75)
    color = _hex_to_rgb(obj.get("fill"))
    bold = obj.get("fontWeight") == "bold"
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(fs)
        run.font.color.rgb = RGBColor(*color)
        run.font.bold = bold


def _add_image_shape(slide, obj):
    src = obj.get("src", "")
    if not src:
        return
    if src.startswith("data:image"):
        data = base64.b64decode(src.split(",", 1)[1])
    else:
        try:
            data = urllib.request.urlopen(src, timeout=15).read()
        except Exception:
            return
    left = int(obj.get("left", 0) * _PX2EMU)
    top = int(obj.get("top", 0) * _PX2EMU)
    width = int((obj.get("width", 100) or 0) * (obj.get("scaleX", 1) or 1) * _PX2EMU)
    height = int((obj.get("height", 100) or 0) * (obj.get("scaleY", 1) or 1) * _PX2EMU)
    if width <= 0 or height <= 0:
        return
    try:
        slide.shapes.add_picture(io.BytesIO(data), left, top, width, height)
    except Exception:
        return

